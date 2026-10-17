from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader


class DocumentReader:

    SUPPORTED_EXTENSIONS = {
        ".pdf",
        ".docx",
        ".xlsx",
        ".txt",
    }

    def read(self, file_path: str) -> str:
        """
        Read a supported candidate document and return its textual content.

        Supported formats:
        - PDF
        - DOCX
        - XLSX
        - TXT
        """

        path = Path(file_path).resolve()

        if not path.exists():
            raise FileNotFoundError(
                f"Document does not exist: {path}"
            )

        if not path.is_file():
            raise ValueError(
                f"Path is not a file: {path}"
            )

        suffix = path.suffix.lower()

        if suffix not in self.SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file format: {suffix}"
            )

        if suffix == ".pdf":
            content = self._read_pdf(path)

        elif suffix == ".docx":
            content = self._read_docx(path)

        elif suffix == ".xlsx":
            content = self._read_xlsx(path)

        elif suffix == ".txt":
            content = self._read_txt(path)

        else:
            raise ValueError(
                f"Unsupported file format: {suffix}"
            )

        if not content.strip():
            raise ValueError(
                f"No readable content found in document: {path.name}"
            )

        return content

    def _read_pdf(self, path: Path) -> str:
        """
        Extract text from a PDF while preserving page boundaries.
        """

        reader = PdfReader(path)

        parts: list[str] = []

        for page_number, page in enumerate(
            reader.pages,
            start=1,
        ):
            text = page.extract_text() or ""

            parts.append(
                f"[Page {page_number}]\n{text}"
            )

        return "\n\n".join(parts)

    def _read_docx(self, path: Path) -> str:
        """
        Extract paragraphs and table contents from a DOCX document.
        """

        document = Document(path)

        parts: list[str] = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()

            if text:
                parts.append(text)

        for table_number, table in enumerate(
            document.tables,
            start=1,
        ):
            parts.append(
                f"[Table {table_number}]"
            )

            for row in table.rows:
                cells = [
                    cell.text.strip()
                    for cell in row.cells
                ]

                parts.append(
                    " | ".join(cells)
                )

        return "\n".join(parts)

    def _read_xlsx(self, path: Path) -> str:
        """
        Extract workbook content while preserving sheet names
        and cell coordinates.
        """

        workbook = load_workbook(
            path,
            read_only=True,
            data_only=True,
        )

        parts: list[str] = []

        try:
            for worksheet in workbook.worksheets:

                parts.append(
                    f"[Sheet: {worksheet.title}]"
                )

                for row in worksheet.iter_rows():

                    cells: list[str] = []

                    for cell in row:

                        if cell.value is None:
                            continue

                        cells.append(
                            f"{cell.coordinate}: {cell.value}"
                        )

                    if cells:
                        parts.append(
                            " | ".join(cells)
                        )

        finally:
            workbook.close()

        return "\n".join(parts)

    def _read_txt(self, path: Path) -> str:
        """
        Read a UTF-8 text file.
        """

        return path.read_text(
            encoding="utf-8"
        )